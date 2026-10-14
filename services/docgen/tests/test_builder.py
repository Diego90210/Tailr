"""DOC-01: the python-docx builder is driven entirely by the format_spec."""
from __future__ import annotations

from copy import deepcopy

from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn

from builder import build_docx
from format_specs import DEFAULT_FORMAT_SPEC

URL = "https://github.com/example/repo"


def _spec(**overrides) -> dict:
    spec = deepcopy(DEFAULT_FORMAT_SPEC)
    spec.update(overrides)
    return spec


def _content() -> dict:
    return {
        "header": {"name": "Ada Lovelace", "contactLine": "ada@example.com"},
        "summary": "Backend engineer.",
        "experience": [{
            "company": "Acme", "role": "Engineer", "start": "2020", "end": "2022",
            "bullets": [{"text": "Did a thing.", "factIds": ["f1"]}],
        }],
        "projects": [{
            "name": "repo", "stack": "Go", "url": URL,
            "bullets": [{"text": "Built it.", "factIds": ["f2"]}],
        }],
        "education": ["B.Sc. CS, 2019"],
        "certifications": [],
        "skills": ["Go"],
        "languages": ["English (native)"],
    }


def _paragraph_containing(doc, text: str):
    for p in doc.paragraphs:
        if any(text in r.text for r in p.runs):
            return p
    raise AssertionError(f"no paragraph containing {text!r}")


def _run_containing(doc, text: str):
    for p in doc.paragraphs:
        for r in p.runs:
            if text in r.text:
                return r
    raise AssertionError(f"no run containing {text!r}")


def test_builder_applies_font_sizes_page_and_margins_from_spec():
    spec = _spec(
        font="Georgia", bodySizePt=13, nameSizePt=20, nameBold=False,
        pageSizeIn={"w": 8.0, "h": 10.0},
        marginsIn={"top": 0.4, "bottom": 0.4, "left": 0.6, "right": 0.6},
    )
    doc = build_docx(_content(), spec)
    assert doc.styles["Normal"].font.name == "Georgia"
    assert doc.styles["Normal"].font.size.pt == 13
    name = doc.paragraphs[0].runs[0]
    assert name.font.size.pt == 20
    assert name.bold is False
    sec = doc.sections[0]
    assert round(sec.page_width.inches, 2) == 8.0
    assert round(sec.page_height.inches, 2) == 10.0
    assert round(sec.left_margin.inches, 2) == 0.6


def test_builder_italic_dates_follow_spec():
    italic = build_docx(_content(), _spec(dateItalic=True))
    assert _run_containing(italic, "2020 – 2022").italic is True
    plain = build_docx(_content(), _spec(dateItalic=False))
    assert _run_containing(plain, "2020 – 2022").italic is not True


def test_builder_hyperlink_targets_url_and_uses_spec_style():
    doc = build_docx(_content(), _spec(linkColor="FF0000", linkUnderline=False))
    links = [h for p in doc.paragraphs for h in p._p.findall(qn("w:hyperlink"))]
    assert len(links) == 1
    r_id = links[0].get(qn("r:id"))
    assert doc.part.rels[r_id].target_ref == URL
    assert links[0].findall(".//" + qn("w:t"))[0].text == URL
    color = links[0].findall(".//" + qn("w:color"))
    assert color and color[0].get(qn("w:val")) == "FF0000"
    assert links[0].findall(".//" + qn("w:u")) == []


def test_builder_right_aligns_dates_with_tab_stop_at_spec_position():
    doc = build_docx(_content(), _spec(tabStopIn=6.0, rightAlignedDates=True))
    para = _paragraph_containing(doc, "2020 – 2022")
    stops = list(para.paragraph_format.tab_stops)
    assert any(s.alignment == WD_TAB_ALIGNMENT.RIGHT and round(s.position.inches, 2) == 6.0
               for s in stops)


def test_builder_appends_dates_inline_when_right_alignment_is_off():
    doc = build_docx(_content(), _spec(rightAlignedDates=False))
    para = _paragraph_containing(doc, "2020 – 2022")
    assert "\t" not in para.text
    assert all(s.alignment != WD_TAB_ALIGNMENT.RIGHT for s in para.paragraph_format.tab_stops)


def test_builder_headings_follow_locale():
    text_en = [p.text for p in build_docx(_content(), _spec(), locale="en").paragraphs]
    text_es = [p.text for p in build_docx(_content(), _spec(), locale="es").paragraphs]
    assert "EXPERIENCE" in text_en and "EXPERIENCIA" not in text_en
    assert "EXPERIENCIA" in text_es and "EXPERIENCE" not in text_es


def test_builder_follows_section_order():
    spec = _spec(sectionOrder=["skills", "summary", "experience"])
    text = [p.text for p in build_docx(_content(), spec).paragraphs]
    assert text.index("SKILLS") < text.index("PROFESSIONAL PROFILE") < text.index("EXPERIENCE")


def test_builder_output_is_deterministic():
    spec = _spec()
    assert build_docx(_content(), spec).element.xml == build_docx(_content(), spec).element.xml
